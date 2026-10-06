import base64
import io
from datetime import date

from openpyxl import load_workbook

from odoo import _, fields, models
from odoo.exceptions import UserError
from odoo.modules.module import get_module_resource


class PrimetechDsfGenerationWizard(models.TransientModel):
    _name = "primetech.dsf.generation.wizard"
    _description = "Génération DSF DGI"

    company_id = fields.Many2one("res.company", string="Société", required=True, default=lambda self: self.env.company)
    # Do not alter the legacy integer ``fiscal_year`` database column.  Older
    # versions of this wizard used it, and PostgreSQL cannot cast it directly
    # to a date during an Odoo upgrade.
    fiscal_year_date = fields.Date(
        string="Date de clôture de l'exercice",
        required=True,
        default=lambda self: fields.Date.context_today(self).replace(month=12, day=31),
        help="Sélectionnez une date de l'exercice à déclarer ; son année est utilisée pour la DSF.",
    )
    file_data = fields.Binary(readonly=True)
    file_name = fields.Char(readonly=True)

    @staticmethod
    def _set_excel_value(sheet, address, value):
        """Write to the anchor cell when the template cell is merged."""
        for merged_range in sheet.merged_cells.ranges:
            if address in merged_range:
                sheet.cell(merged_range.min_row, merged_range.min_col).value = value
                return
        sheet[address] = value

    def _set_document_headers(self, workbook):
        """Fill common company/exercise information without touching styles."""
        company = self.company_id
        closing = "31/12/%s" % self.fiscal_year_date.year
        designation = "Désignation entité : %s    Exercice clos le %s" % (company.name or "", closing)
        identification = "Numéro d’identification : %s    Durée (en mois) : 12" % (company.vat or "")
        for sheet in workbook.worksheets:
            # The official DSF template has a huge formatted grid.  Iterating
            # rows would create every empty cell and exhaust the worker memory.
            # Work only on cells that already exist in the workbook.
            for row in (sheet._cells.values(),):
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str):
                        continue
                    normalized = value.replace("\xa0", " ")
                    if "Désignation entité" in normalized:
                        cell.value = designation
                    elif "Numéro d’identification" in normalized:
                        cell.value = identification
                    elif value == "COMPTE RESULTAT AU 31 DECEMBRE ______":
                        cell.value = "COMPTE DE RÉSULTAT AU 31 DÉCEMBRE %s" % self.fiscal_year_date.year
                    elif value == "BILAN AU 31 DECEMBRE N":
                        cell.value = "BILAN AU 31 DÉCEMBRE %s" % self.fiscal_year_date.year

        cover = workbook["PAGE DE GARDE"]
        self._set_excel_value(cover, "B10", "CENTRE DE DEPOT DE: %s" % (company.city or ""))
        self._set_excel_value(cover, "C18", closing)
        self._set_excel_value(cover, "A25", "DENOMINATION SOCIALE: %s" % (company.name or ""))
        self._set_excel_value(
            cover,
            "A31",
            "ADRESSE COMPLETE: %s" % ", ".join(
                filter(None, [company.street, company.street2, company.zip, company.city, company.country_id.name])
            ),
        )
        self._set_excel_value(cover, "A33", "N° D'IDENTIFICATION FISCALE: %s" % (company.vat or ""))
        return
        cover["A33"] = "N° D'IDENTIFICATION FISCALE: %s" % (company.vat or "")

    def action_generate_dsf(self):
        self.ensure_one()
        if not self.fiscal_year_date or self.fiscal_year_date.year < 2000 or self.fiscal_year_date.year > 2100:
            raise UserError(_("Veuillez saisir une année d'exercice valide."))
        template_path = get_module_resource("primetech_tax_automation", "data", "dsf_normal_dgiformat.xlsx")
        if not template_path:
            raise UserError(_("Le modèle DSF officiel est introuvable dans le module."))
        workbook = load_workbook(template_path)
        self._set_document_headers(workbook)
        date_from = date(self.fiscal_year_date.year, 1, 1)
        date_to = date(self.fiscal_year_date.year, 12, 31)
        mappings = self.env["primetech.dsf.mapping"].search([
            ("company_id", "=", self.company_id.id), ("active", "=", True),
        ])
        if not mappings:
            raise UserError(_(
                "Aucune correspondance DSF n'est configurée. Ajoutez au moins une ligne dans Configuration › Correspondances DSF."
            ))
        for mapping in mappings:
            if mapping.sheet_name not in workbook.sheetnames:
                raise UserError(_("La feuille DSF « %s » n'existe pas dans le modèle officiel.") % mapping.sheet_name)
            try:
                workbook[mapping.sheet_name][mapping.cell_address].value = round(
                    mapping._amount_for_period(date_from, date_to), 0
                )
            except ValueError:
                raise UserError(_("Cellule Excel invalide : %s (%s).") % (mapping.cell_address, mapping.name))
        output = io.BytesIO()
        workbook.save(output)
        filename = "DSF_%s_%s.xlsx" % (self.company_id.name.replace("/", "-"), self.fiscal_year_date.year)
        self.write({"file_data": base64.b64encode(output.getvalue()), "file_name": filename})
        return {
            "type": "ir.actions.act_url",
            "url": "/web/content?model=primetech.dsf.generation.wizard&id=%s&field=file_data&filename_field=file_name&download=true" % self.id,
            "target": "self",
        }
