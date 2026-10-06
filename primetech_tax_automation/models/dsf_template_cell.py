"""Selectable cells of the official DGI DSF workbook."""

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from odoo import api, fields, models
from odoo.modules.module import get_module_resource


class PrimetechDsfTemplateCell(models.Model):
    _name = "primetech.dsf.template.cell"
    _description = "Cellule disponible du modèle DSF"
    _order = "sheet_name, cell_address"

    name = fields.Char(required=True, readonly=True)
    sheet_name = fields.Char(required=True, readonly=True, index=True)
    cell_address = fields.Char(required=True, readonly=True, index=True)

    _sql_constraints = [
        ("dsf_template_cell_unique", "unique(sheet_name, cell_address)", "Cette cellule DSF existe déjà."),
    ]

    @api.model
    def ensure_sheet_cells(self, sheet_name):
        """Index only numeric, writable cells of the selected DSF sheet."""
        if not sheet_name:
            return self.browse()
        existing = self.search([("sheet_name", "=", sheet_name)])
        if existing:
            return existing

        template_path = get_module_resource("primetech_tax_automation", "data", "dsf_normal_dgiformat.xlsx")
        workbook = load_workbook(template_path, read_only=False, data_only=False)
        if sheet_name not in workbook.sheetnames:
            return self.browse()
        sheet = workbook[sheet_name]
        used_rows = {cell.row for cell in sheet._cells.values() if cell.value not in (None, "")}
        values = []
        for cell in sheet._cells.values():
            if (
                isinstance(cell, MergedCell)
                or cell.value is not None
                or cell.row not in used_rows
                or "#,##0" not in (cell.number_format or "")
            ):
                continue
            label = ""
            for column in range(cell.column - 1, 0, -1):
                neighbour = sheet._cells.get((cell.row, column))
                if neighbour and isinstance(neighbour.value, str) and neighbour.value.strip():
                    label = " ".join(neighbour.value.split())[:100]
                    break
            display_name = cell.coordinate if not label else "%s — %s" % (cell.coordinate, label)
            values.append({
                "name": display_name,
                "sheet_name": sheet_name,
                "cell_address": cell.coordinate,
            })
        return self.create(values)
